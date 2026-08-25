#!/usr/bin/env python3
"""Generate 50 state crime collection pages from official FBI and Census files.

Inputs are deliberately local snapshots. The script never downloads data and never
uses CrimeByCity. It ranks Census-matched places by the 2024 FBI Table 8 reported
violent + property crime rate among agencies covering at least 10,000 residents.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import yaml


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / ".tmp/crime-fbi/CIUS_Table_8_Offenses_Known_to_Law_Enforcement_by_State_by_City_2024.xlsx"
DEFAULT_GAZETTEER = ROOT / ".tmp/crime-fbi/2024_Gaz_place_national.txt"
DEFAULT_OUTPUT = ROOT / "Content/collections/crime"
DEFAULT_REPORT = ROOT / ".tmp/crime-fbi/crime-generation-summary.json"
DEFAULT_PHOTO_MANIFEST = DEFAULT_OUTPUT / "image-credits.json"
DEFAULT_CITY_PHOTO_MANIFEST = DEFAULT_OUTPUT / "city-image-credits.json"

STATES = [
    ("Alabama", "AL"), ("Alaska", "AK"), ("Arizona", "AZ"), ("Arkansas", "AR"),
    ("California", "CA"), ("Colorado", "CO"), ("Connecticut", "CT"), ("Delaware", "DE"),
    ("Florida", "FL"), ("Georgia", "GA"), ("Hawaii", "HI"), ("Idaho", "ID"),
    ("Illinois", "IL"), ("Indiana", "IN"), ("Iowa", "IA"), ("Kansas", "KS"),
    ("Kentucky", "KY"), ("Louisiana", "LA"), ("Maine", "ME"), ("Maryland", "MD"),
    ("Massachusetts", "MA"), ("Michigan", "MI"), ("Minnesota", "MN"), ("Mississippi", "MS"),
    ("Missouri", "MO"), ("Montana", "MT"), ("Nebraska", "NE"), ("Nevada", "NV"),
    ("New Hampshire", "NH"), ("New Jersey", "NJ"), ("New Mexico", "NM"), ("New York", "NY"),
    ("North Carolina", "NC"), ("North Dakota", "ND"), ("Ohio", "OH"), ("Oklahoma", "OK"),
    ("Oregon", "OR"), ("Pennsylvania", "PA"), ("Rhode Island", "RI"),
    ("South Carolina", "SC"), ("South Dakota", "SD"), ("Tennessee", "TN"), ("Texas", "TX"),
    ("Utah", "UT"), ("Vermont", "VT"), ("Virginia", "VA"), ("Washington", "WA"),
    ("West Virginia", "WV"), ("Wisconsin", "WI"), ("Wyoming", "WY"),
]
STATE_BY_UPPER = {name.upper(): (name, postal) for name, postal in STATES}
STATE_BY_POSTAL = {postal: name for name, postal in STATES}
POSTAL_BY_STATE = {name: postal for name, postal in STATES}

# These are agency labels that do not directly equal a Census place label.
# Values are normalized Gazetteer place names, not hand-entered coordinates.
PLACE_ALIASES = {
    ("FL", "lake worth"): "lake worth beach",
    ("HI", "honolulu"): "urban honolulu",
    ("ID", "boise"): "boise city",
    ("KY", "louisville metro"): "louisville jefferson county",
    ("NV", "las vegas metropolitan police department"): "las vegas",
    ("NC", "charlotte mecklenburg"): "charlotte",
    ("TN", "metropolitan nashville police department"): "nashville davidson",
    ("UT", "west valley"): "west valley city",
    ("MA", "west springfield"): "west springfield town",
}

DISPLAY_NAMES = {
    ("KY", "louisville metro"): "Louisville",
    ("NV", "las vegas metropolitan police department"): "Las Vegas",
    ("TN", "metropolitan nashville police department"): "Nashville",
}

FBI_SOURCE = "https://cde.ucr.cjis.gov/LATEST/"
CENSUS_SOURCE = "https://www2.census.gov/geo/docs/maps-data/data/gazetteer/2024_Gazetteer/2024_Gaz_place_national.zip"


class LiteralString(str):
    pass


class Dumper(yaml.SafeDumper):
    pass


Dumper.add_representer(
    LiteralString,
    lambda dumper, value: dumper.represent_scalar("tag:yaml.org,2002:str", value, style="|"),
)


@dataclass
class Place:
    postal: str
    name: str
    latitude: float
    longitude: float
    priority: int


@dataclass
class CrimeRow:
    state: str
    postal: str
    city: str
    population: int
    violent: int
    murder: int
    rape: int
    robbery: int
    assault: int
    property: int
    burglary: int
    larceny: int
    vehicle_theft: int
    arson: int | None
    place: Place | None = None

    @property
    def total(self) -> int:
        return self.violent + self.property

    def rate(self, count: int) -> float:
        return round(count / self.population * 100_000, 1)

    @property
    def total_rate(self) -> float:
        return self.rate(self.total)

    @property
    def violent_rate(self) -> float:
        return self.rate(self.violent)

    @property
    def property_rate(self) -> float:
        return self.rate(self.property)

    @property
    def murder_rate(self) -> float:
        return self.rate(self.murder)

    @property
    def display_city(self) -> str:
        return DISPLAY_NAMES.get((self.postal, normalize_name(self.city)), self.city)


def slugify(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def normalize_name(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    value = value.lower().replace("&", " and ").replace("/", " ")
    value = re.sub(r"\(balance\)", " ", value)
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def gazetteer_base_name(value: str) -> str:
    value = normalize_name(value)
    suffixes = [
        "city and borough", "metropolitan government", "metro government",
        "consolidated government", "unified government", "metropolitan municipality",
        "municipality", "borough", "village", "town", "city", "cdp",
    ]
    for suffix in suffixes:
        if value.endswith(" " + suffix):
            return value[: -(len(suffix) + 1)].strip()
    return value


def place_priority(name: str, funcstat: str) -> int:
    """Prefer active incorporated places over same-named CDPs."""
    normalized = normalize_name(name)
    if funcstat == "A" and not normalized.endswith(" cdp"):
        return 0
    if normalized.endswith(" cdp"):
        return 3
    return 2


def normalize_state(raw: Any) -> tuple[str, str] | None:
    cleaned = re.sub(r"[^A-Z ]", "", str(raw).upper())
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if cleaned == "DISTRICT OF COLUMBIA":
        return None
    return STATE_BY_UPPER.get(cleaned)


def number(value: Any, *, allow_none: bool = False) -> int | None:
    if value is None or value == "":
        return None if allow_none else 0
    if not isinstance(value, (int, float)):
        raise ValueError(f"Expected numeric XLSX cell, got {value!r}")
    return int(value)


def read_fbi(path: Path, minimum_population: int) -> tuple[list[CrimeRow], list[str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["24tbl08"] if "24tbl08" in workbook.sheetnames else workbook.active
    headers = [str(cell.value or "").replace("\n", " ").strip() for cell in sheet[4][:13]]
    expected = ["State", "City", "Population", "Violent crime"]
    if headers[:4] != expected:
        raise ValueError(f"Unexpected FBI header row: {headers}")

    rows: list[CrimeRow] = []
    warnings: list[str] = []
    for values in sheet.iter_rows(min_row=5, max_col=13, values_only=True):
        state_info = normalize_state(values[0])
        if state_info is None:
            continue
        state, postal = state_info
        city = str(values[1] or "").strip()
        population = number(values[2]) or 0
        if not city or population < minimum_population:
            continue
        row = CrimeRow(
            state=state, postal=postal, city=city, population=population,
            violent=number(values[3]) or 0, murder=number(values[4]) or 0,
            rape=number(values[5]) or 0, robbery=number(values[6]) or 0,
            assault=number(values[7]) or 0, property=number(values[8]) or 0,
            burglary=number(values[9]) or 0, larceny=number(values[10]) or 0,
            vehicle_theft=number(values[11]) or 0, arson=number(values[12], allow_none=True),
        )
        if row.violent != row.murder + row.rape + row.robbery + row.assault:
            warnings.append(f"{state}/{city}: violent total does not equal components")
        if row.property != row.burglary + row.larceny + row.vehicle_theft:
            warnings.append(f"{state}/{city}: property total does not equal components")
        rows.append(row)
    return rows, warnings


def read_gazetteer(path: Path) -> tuple[dict[tuple[str, str], Place], list[str]]:
    candidates: dict[tuple[str, str], list[Place]] = defaultdict(list)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for raw in csv.DictReader(handle, delimiter="\t"):
            row = {str(key).strip(): str(value).strip() for key, value in raw.items() if key is not None}
            postal = row["USPS"]
            if postal not in STATE_BY_POSTAL:
                continue
            place = Place(
                postal, row["NAME"], float(row["INTPTLAT"]), float(row["INTPTLONG"]),
                place_priority(row["NAME"], row.get("FUNCSTAT", "")),
            )
            for key in {normalize_name(place.name), gazetteer_base_name(place.name)}:
                candidates[(postal, key)].append(place)

    lookup: dict[tuple[str, str], Place] = {}
    ambiguous: list[str] = []
    for key, places in candidates.items():
        unique = {(p.name, p.latitude, p.longitude): p for p in places}
        best_priority = min(place.priority for place in unique.values())
        best = [place for place in unique.values() if place.priority == best_priority]
        if len(best) == 1:
            lookup[key] = best[0]
        else:
            ambiguous.append(f"{key[0]}/{key[1]}: {', '.join(sorted(p.name for p in best))}")
    return lookup, ambiguous


def attach_places(rows: list[CrimeRow], places: dict[tuple[str, str], Place]) -> list[dict[str, Any]]:
    unmatched: list[dict[str, Any]] = []
    for row in rows:
        raw_key = normalize_name(row.city)
        target = PLACE_ALIASES.get((row.postal, raw_key), raw_key)
        row.place = places.get((row.postal, target))
        if row.place is None:
            unmatched.append({
                "state": row.state, "postal_code": row.postal, "agency_city": row.city,
                "coverage_population": row.population, "total_crime_rate": row.total_rate,
            })
    return unmatched


def fmt_rate(value: float) -> str:
    return f"{value:,.1f}"


def page_for_state(
    state: str,
    postal: str,
    rows: list[CrimeRow],
    generated_on: str,
    photo: dict[str, Any] | None = None,
    city_photos: dict[str, dict[str, Any]] | None = None,
) -> dict[str, Any]:
    state_slug = slugify(state)
    ranked = sorted(rows, key=lambda row: (-row.total_rate, -row.population, row.display_city))[:10]
    if not ranked:
        raise ValueError(f"No Census-matched FBI rows for {state}")
    leader = ranked[0]
    count = len(ranked)
    leader_property_share = (leader.property / leader.total * 100) if leader.total else 0.0
    violent_rate_leader = max(ranked, key=lambda row: row.violent_rate)
    smallest_coverage = min(ranked, key=lambda row: row.population)
    list_label = "10 cities" if count == 10 else f"{count} qualifying {'city' if count == 1 else 'cities'}"
    population_label = "a coverage population" if count == 1 else "coverage populations"
    seo_title = f"Most Dangerous {'City' if count == 1 else 'Cities'} in {state} | 2024 Data"
    seo_description = (
        f"{leader.display_city} leads qualifying {state} city agencies in 2024 with "
        f"{fmt_rate(leader.total_rate)} reported violent and property crimes per 100,000."
    )
    if len(seo_title) > 58 or len(seo_description) > 152:
        raise ValueError(f"SEO limit exceeded for {state}: {len(seo_title)}/{len(seo_description)}")

    rows_yaml = []
    markers = []
    highlights = []
    for rank, row in enumerate(ranked, 1):
        assert row.place is not None
        city_name = row.display_city
        city_slug = slugify(city_name)
        dominant = "property crime" if row.property >= row.violent else "violent crime"
        dominant_rate = row.property_rate if row.property >= row.violent else row.violent_rate
        rows_yaml.append({
            "rank": rank, "city": city_name, "city_slug": city_slug,
            "coverage_population": row.population, "total_crimes": row.total,
            "total_crime_rate": row.total_rate, "violent_crimes": row.violent,
            "violent_crime_rate": row.violent_rate, "property_crimes": row.property,
            "property_crime_rate": row.property_rate, "murders": row.murder,
            "murder_rate": row.murder_rate,
        })
        markers.append({
            "rank": rank, "slug": city_slug,
            "anchor": f"highest-reported-crime-rates-{city_slug}", "name": city_name,
            "latitude": row.place.latitude, "longitude": row.place.longitude,
            "total_crime_rate": row.total_rate, "violent_crime_rate": row.violent_rate,
            "property_crime_rate": row.property_rate, "coverage_population": row.population,
            "popup_text": f"{dominant.capitalize()} is the larger category at {fmt_rate(dominant_rate)} per 100,000.",
        })
        highlight = {
            "name": city_name, "state": state, "status": "2024 FBI data",
            "description": LiteralString(
                f"**Rank:** #{rank}\n"
                f"**Total crime rate:** {fmt_rate(row.total_rate)} per 100,000\n"
                f"**Coverage population:** {row.population:,}\n\n"
                f"{city_name} reported {row.total:,} violent and property crimes in 2024. "
                f"The violent crime rate was {fmt_rate(row.violent_rate)}, while the property crime rate was "
                f"{fmt_rate(row.property_rate)} per 100,000.\n\n"
                f"{dominant.capitalize()} accounts for the larger share of the combined rate, which places {city_name} "
                f"at #{rank} among qualifying {state} city agencies matched to Census places."
            ),
        }
        city_photo = (city_photos or {}).get(f"{state_slug}/{city_slug}")
        if city_photo:
            city_image = str(city_photo.get("local_path") or "").strip()
            if city_image:
                highlight["image"] = city_image
        highlights.append(highlight)

    noun = "city" if count == 1 else "cities"
    if count == 1:
        ranking_reason_paragraphs = [
            f"{leader.display_city} is the only qualifying {state} city agency in this comparison. It reported "
            f"{leader.total:,} violent and property crimes, producing a combined rate of "
            f"{fmt_rate(leader.total_rate)} per 100,000 residents covered.",
            f"Property crime contributes {fmt_rate(leader.property_rate)} per 100,000, or "
            f"{leader_property_share:.1f}% of the combined rate. Violent crime contributes the remaining "
            f"{fmt_rate(leader.violent_rate)} per 100,000. The first-place label therefore reflects the only "
            "eligible agency in the dataset rather than a comparison with every community in the state.",
        ]
        standout_facts = [
            f"{leader.display_city} reported {leader.total:,} combined violent and property offenses for an FBI "
            f"coverage population of {leader.population:,}.",
            f"Property offenses account for {leader_property_share:.1f}% of its combined reported crime rate; "
            f"violent offenses account for {100 - leader_property_share:.1f}%.",
            f"No second qualifying {state} city is available in this matched dataset, so a statewide rate gap "
            "cannot be calculated.",
        ]
    else:
        runner_up = ranked[1]
        rate_gap = leader.total_rate - runner_up.total_rate
        one_offense_effect = 100_000 / smallest_coverage.population
        rank_by_city = {row.display_city: rank for rank, row in enumerate(ranked, 1)}
        incident_leader = max(ranked, key=lambda row: row.total)
        property_rate_leader = max(ranked, key=lambda row: row.property_rate)
        property_shares = [
            (row.property / row.total * 100 if row.total else 0.0, row)
            for row in ranked
        ]
        highest_property_share, highest_property_share_city = max(property_shares, key=lambda item: item[0])
        lowest_property_share, lowest_property_share_city = min(property_shares, key=lambda item: item[0])
        adjacent_gaps = [
            (left.total_rate - right.total_rate, left, right)
            for left, right in zip(ranked, ranked[1:])
        ]
        largest_gap, gap_above_city, gap_below_city = max(adjacent_gaps, key=lambda item: item[0])
        ranking_reason_paragraphs = [
            f"{leader.display_city} ranks first because its combined 2024 reported crime rate is "
            f"{fmt_rate(leader.total_rate)} per 100,000, which is {fmt_rate(rate_gap)} points above "
            f"{runner_up.display_city}. Property crime supplies {leader_property_share:.1f}% of "
            f"{leader.display_city}'s total, while violent crime contributes {fmt_rate(leader.violent_rate)} "
            "points per 100,000.",
            f"The total-crime order is not automatically the violent-crime order. "
            f"{violent_rate_leader.display_city} has the highest violent-crime rate in this group at "
            f"{fmt_rate(violent_rate_leader.violent_rate)} per 100,000, while the published ranking combines "
            "violent and property offenses.",
            f"Coverage size also affects how quickly a rate moves. {smallest_coverage.display_city} has the "
            f"smallest FBI coverage population in this list at {smallest_coverage.population:,}; one additional "
            f"reported offense changes its rate by about {one_offense_effect:.1f} per 100,000. This mathematical "
            "effect explains volatility in smaller-city rates, not the underlying social causes of crime.",
        ]
        standout_facts = []
        if incident_leader.display_city == leader.display_city:
            standout_facts.append(
                f"{leader.display_city} leads both the per-capita ranking and the raw combined offense count "
                f"among the included {state} cities, with {leader.total:,} reported offenses."
            )
        else:
            standout_facts.append(
                f"{incident_leader.display_city} reports the largest raw combined offense count in this group "
                f"({incident_leader.total:,}), but {leader.display_city} ranks first after population adjustment."
            )
        if violent_rate_leader.display_city == leader.display_city:
            standout_facts.append(
                f"{leader.display_city} also has the group's highest violent-crime rate, so its first-place total "
                "is not driven by property crime alone."
            )
        else:
            standout_facts.append(
                f"{violent_rate_leader.display_city}, ranked #{rank_by_city[violent_rate_leader.display_city]} "
                f"overall, has the highest violent-crime rate; {leader.display_city}'s #1 position comes from the "
                "combined violent-and-property measure."
            )
        standout_facts.append(
            f"The list's largest adjacent gap is {fmt_rate(largest_gap)} rate points between "
            f"{gap_above_city.display_city} and {gap_below_city.display_city}."
        )
        standout_facts.append(
            f"Property crime's share ranges from {lowest_property_share:.1f}% in "
            f"{lowest_property_share_city.display_city} to {highest_property_share:.1f}% in "
            f"{highest_property_share_city.display_city}."
        )
        if property_rate_leader.display_city not in {
            leader.display_city,
            violent_rate_leader.display_city,
            incident_leader.display_city,
        }:
            standout_facts.append(
                f"{property_rate_leader.display_city} has the highest property-crime rate in the group even though "
                f"it ranks #{rank_by_city[property_rate_leader.display_city]} on the combined list."
            )
    page = {
        "type": "collection", "slug": state_slug, "category": "crime",
        "url": f"/most-dangerous-cities/{state_slug}",
        "auto_link_phrases": [
            f"most dangerous cities in {state}", f"highest crime cities in {state}",
            f"{state} city crime rates",
        ],
        "author": "USA Symbol Team", "date_published": generated_on, "date_modified": generated_on,
        "seo": {"title": seo_title, "description": seo_description},
        "page": {
            "h1": f"Most Dangerous {'City' if count == 1 else 'Cities'} in {state}",
            "intro_title": f"Highest Reported City Crime Rates in {state}",
            "quick_answer": [
                f"{leader.display_city} has the highest 2024 reported total crime rate among qualifying {state} city agencies matched to Census places at {fmt_rate(leader.total_rate)} per 100,000 residents covered.",
                f"Its rate combines {fmt_rate(leader.violent_rate)} violent crimes and {fmt_rate(leader.property_rate)} property crimes per 100,000. The table compares {list_label} with {population_label} of at least 10,000.",
                "These are city-agency reporting rates, not neighborhood risk estimates. Reporting practices, agency boundaries, and the number of crimes reported to police affect the comparison.",
            ],
            "methodology": (
                "Cities are ranked by total reported crime rate, calculated as FBI Table 8 violent crime plus "
                "property crime divided by the agency coverage population and multiplied by 100,000. Eligibility "
                "requires a 2024 coverage population of at least 10,000 and a matching 2024 Census Gazetteer place; "
                "the table is limited to the top 10. Arson is not included in the FBI property-crime total. Display "
                "names may simplify an FBI agency label, while rates continue to use that agency's coverage population."
            ),
            "sources": [
                {"name": "FBI Crime Data Explorer, 2024 CIUS Table 8", "url": FBI_SOURCE,
                 "description": "Official city-agency population and reported offense counts used for every rate"},
                {"name": "U.S. Census Bureau 2024 Gazetteer, Places", "url": CENSUS_SOURCE,
                 "description": "Official place names and internal-point coordinates used for map markers"},
            ],
        },
        "table": {
            "title": f"Highest Reported Crime {'Rate' if count == 1 else 'Rates'} in {state}, 2024",
            "caption": f"{count} qualifying {state} {noun} ranked by combined violent and property crime rate.",
            "note": "Rates are per 100,000 residents in the reporting agency's coverage population, which may differ from Census place population.",
            "sortable": True, "sort_default": "desc", "metric_default": "total_crime_rate",
            "column_formats": {
                "total_crime_rate": "0.0", "violent_crime_rate": "0.0",
                "property_crime_rate": "0.0", "murder_rate": "0.0",
            },
            "columns": {
                "rank": "Rank", "city": "City", "coverage_population": "Coverage Population",
                "total_crime_rate": "Total Crime Rate", "violent_crime_rate": "Violent Crime Rate",
                "property_crime_rate": "Property Crime Rate", "murder_rate": "Murder Rate",
            },
            "rows": rows_yaml,
        },
        "point_map": {
            "title": f"Map of the {'City' if count == 1 else 'Cities'} With the Highest Reported Crime {'Rate' if count == 1 else 'Rates'} in {state}",
            "image_alt": f"Interactive map showing {count} {state} {'city' if count == 1 else 'cities'} in the 2024 FBI ranking",
            "caption": ("The numbered marker follows the table rank. Select it for rates and a link to the city summary."
                        if count == 1 else
                        "Numbered markers follow the table rank. Select a marker for rates and a link to the city summary."),
            "max_zoom": 9, "markers": markers,
        },
        "sections": [
            {
                "id": "highest-reported-crime-rates", "icon": "fa-solid fa-location-dot",
                "style": "law-cards",
                "title": f"{state} {'City' if count == 1 else 'Cities'} With the Highest Reported Crime {'Rate' if count == 1 else 'Rates'}",
                "paragraphs": [
                    "Each card separates violent and property crime so the category driving the combined rate remains visible. Counts describe offenses reported by the listed city agency during 2024."
                ],
                "highlights": highlights,
            },
            {
                "id": "why-these-cities-rank-highest", "icon": "fa-solid fa-chart-line",
                "title": (
                    f"Why {leader.display_city} Has the Highest Reported Crime Rate in {state}"
                    if count == 1 else
                    f"Why These {state} Cities Rank Highest for Reported Crime"
                ),
                "paragraphs": ranking_reason_paragraphs,
                "facts": standout_facts,
            },
        ],
        "faq": [
            {"question": f"What city has the highest reported crime rate in {state}?",
             "answer": f"{leader.display_city} ranks first among qualifying city agencies matched to Census places at {fmt_rate(leader.total_rate)} total reported crimes per 100,000 residents covered in 2024."},
            {"question": f"Why does {leader.display_city} rank first for reported crime in {state}?",
             "answer": (
                 f"Its combined violent and property crime rate is {fmt_rate(leader.total_rate)} per 100,000 "
                 f"residents covered. Property crime contributes {fmt_rate(leader.property_rate)} and violent "
                 f"crime contributes {fmt_rate(leader.violent_rate)} per 100,000 to that total."
             )},
            {"question": "What crimes are included in the total crime rate?",
             "answer": "Violent crime includes murder, rape, robbery, and aggravated assault. Property crime includes burglary, larceny-theft, and motor vehicle theft. FBI Table 8 lists arson separately, so it is not added to the property total."},
            {"question": f"Does a high city crime rate mean every neighborhood in {state} is dangerous?",
             "answer": "No. FBI Table 8 reports at the law-enforcement agency level. Crime can vary sharply within a city, and the dataset does not provide neighborhood-level risk."},
            {"question": "Why can coverage population differ from city population?",
             "answer": "A police agency's jurisdiction may not match Census place boundaries exactly. The calculation uses the FBI coverage population because that is the population associated with the reported offenses."},
        ],
        "related": [
            {"title": "Crime Rate by State", "url": "/rankings/crime/crime-rate-by-state"},
            {"title": "Safest Cities by State", "url": "/rankings/crime/safest-cities-by-state"},
            {"title": "Most Dangerous City in Each State", "url": "/rankings/crime/most-dangerous-cities-by-state"},
        ],
    }

    if photo:
        provider = str(photo.get("provider") or "Wikimedia Commons").strip()
        author = str(photo.get("author") or "Wikimedia Commons contributor").strip()
        license_name = str(photo.get("license") or "license listed in the image credits").strip()
        local_path = str(photo.get("local_path") or "").strip()
        if local_path:
            photo_alt = f"City view of {leader.display_city}, {state}"
            photo_caption = (
                f"{leader.display_city}, {state}. Photo by {author} via {provider} "
                f"({license_name}); full source details are recorded in the image credits manifest."
            )
            page["hero_image"] = local_path
            page["hero_image_alt"] = photo_alt
            page["hero_image_caption"] = photo_caption
            page["visual_assets"] = [{
                "id": f"{state_slug}-{slugify(leader.display_city)}-city-view",
                "src": local_path,
                "alt": photo_alt,
                "caption": photo_caption,
                "layout": "full-width-wide",
                "section": "highest-reported-crime-rates",
            }]

    return page


def validate_page(page: dict[str, Any], state: str) -> list[str]:
    errors: list[str] = []
    if len(page["seo"]["title"]) > 58:
        errors.append(f"{state}: SEO title exceeds 58 characters")
    if len(page["seo"]["description"]) > 152:
        errors.append(f"{state}: SEO description exceeds 152 characters")
    rows = page["table"]["rows"]
    markers = page["point_map"]["markers"]
    cards = page["sections"][0]["highlights"]
    if not (1 <= len(rows) <= 10) or len(rows) != len(markers) or len(rows) != len(cards):
        errors.append(f"{state}: inconsistent table/marker/card counts")
    if [row["rank"] for row in rows] != list(range(1, len(rows) + 1)):
        errors.append(f"{state}: invalid ranks")
    if any(rows[i]["total_crime_rate"] < rows[i + 1]["total_crime_rate"] for i in range(len(rows) - 1)):
        errors.append(f"{state}: table is not descending")
    if any(marker["coverage_population"] < 10_000 for marker in markers):
        errors.append(f"{state}: marker below population threshold")
    if page["url"] != f"/most-dangerous-cities/{slugify(state)}":
        errors.append(f"{state}: invalid canonical route")
    if len(page["sections"]) != 2 or page["sections"][1]["id"] != "why-these-cities-rank-highest":
        errors.append(f"{state}: missing ranking-reasons section")
    return errors


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--gazetteer", type=Path, default=DEFAULT_GAZETTEER)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--photo-manifest", type=Path, default=DEFAULT_PHOTO_MANIFEST)
    parser.add_argument("--city-photo-manifest", type=Path, default=DEFAULT_CITY_PHOTO_MANIFEST)
    parser.add_argument("--minimum-population", type=int, default=10_000)
    parser.add_argument("--date", default=date.today().isoformat())
    parser.add_argument("--validate-only", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.xlsx.is_file() or not args.gazetteer.is_file():
        raise FileNotFoundError(f"Missing input: xlsx={args.xlsx} gazetteer={args.gazetteer}")
    fbi_rows, source_warnings = read_fbi(args.xlsx, args.minimum_population)
    places, ambiguous = read_gazetteer(args.gazetteer)
    unmatched = attach_places(fbi_rows, places)
    photo_manifest: dict[str, dict[str, Any]] = {}
    if args.photo_manifest.is_file():
        manifest_data = json.loads(args.photo_manifest.read_text(encoding="utf-8"))
        if isinstance(manifest_data, dict) and isinstance(manifest_data.get("photos"), dict):
            manifest_data = manifest_data["photos"]
        if not isinstance(manifest_data, dict):
            raise ValueError(f"Photo manifest must be an object keyed by state slug: {args.photo_manifest}")
        photo_manifest = {
            str(key): value for key, value in manifest_data.items() if isinstance(value, dict)
        }
    city_photo_manifest: dict[str, dict[str, Any]] = {}
    if args.city_photo_manifest.is_file():
        city_manifest_data = json.loads(args.city_photo_manifest.read_text(encoding="utf-8"))
        if isinstance(city_manifest_data, dict) and isinstance(city_manifest_data.get("photos"), dict):
            city_manifest_data = city_manifest_data["photos"]
        if not isinstance(city_manifest_data, dict):
            raise ValueError(
                f"City photo manifest must be an object keyed by state/city slug: {args.city_photo_manifest}"
            )
        city_photo_manifest = {
            str(key): value for key, value in city_manifest_data.items() if isinstance(value, dict)
        }
    matched_by_state: dict[str, list[CrimeRow]] = defaultdict(list)
    for row in fbi_rows:
        if row.place is not None:
            matched_by_state[row.state].append(row)

    pages: dict[str, dict[str, Any]] = {}
    validation_errors: list[str] = []
    for state, postal in STATES:
        page = page_for_state(
            state,
            postal,
            matched_by_state[state],
            args.date,
            photo_manifest.get(slugify(state)),
            city_photo_manifest,
        )
        pages[state] = page
        validation_errors.extend(validate_page(page, state))

    if len(pages) != 50 or "District of Columbia" in pages:
        validation_errors.append("Output must contain exactly 50 states and exclude DC")
    report = {
        "generated_on": args.date, "xlsx": str(args.xlsx), "gazetteer": str(args.gazetteer),
        "minimum_population": args.minimum_population, "fbi_eligible_rows": len(fbi_rows),
        "photo_manifest": str(args.photo_manifest), "photos_attached": sum(
            1 for page in pages.values() if page.get("visual_assets")
        ),
        "city_photo_manifest": str(args.city_photo_manifest), "city_card_photos_attached": sum(
            1
            for page in pages.values()
            for highlight in page["sections"][0]["highlights"]
            if highlight.get("image")
        ),
        "census_matched_rows": sum(len(value) for value in matched_by_state.values()),
        "unmatched_count": len(unmatched), "unmatched": unmatched,
        "ambiguous_gazetteer_keys_ignored": ambiguous, "source_warnings": source_warnings,
        "validation_errors": validation_errors,
        "states": {state: {"eligible_matched": len(matched_by_state[state]),
                            "published": len(pages[state]["table"]["rows"])} for state, _ in STATES},
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    if validation_errors:
        print("Validation failed:\n" + "\n".join(validation_errors), file=sys.stderr)
        return 1
    if not args.validate_only:
        args.output.mkdir(parents=True, exist_ok=True)
        for state, page in pages.items():
            target = args.output / f"{slugify(state)}.yml"
            target.write_text(yaml.dump(page, Dumper=Dumper, sort_keys=False, allow_unicode=True, width=110), encoding="utf-8")
    print(json.dumps({"states": 50, "files_written": 0 if args.validate_only else 50,
                      "eligible": len(fbi_rows), "matched": report["census_matched_rows"],
                      "unmatched": len(unmatched), "report": str(args.report)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
